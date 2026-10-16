"""
Embedded video support.

Typst has no video element, so a deck marks each video with a poster image
plus a ``<typ2pptx-video>`` metadata entry (see ``examples/video/video.typ``)::

    (src: "videos/demo.mp4", poster: "videos/demo.jpg",
     page: 2, x: 279.2, y: 160.5, width: 283.5, height: 159.4)

``page`` is the physical (1-based) page and ``x``/``y``/``width``/``height``
are absolute page coordinates in pt, i.e. the typst.ts SVG coordinate space.

After the slides are converted, each poster picture that typ2pptx placed from
the SVG is replaced, at the same z-order position, by a movie shape that embeds
the video file. A metadata entry without a matching poster picture on its page
is skipped: Touying's ``#pause`` / ``#only`` hide content but keep its
metadata, so this is how videos on not-yet-revealed subslides are left out.
"""

import json
import mimetypes
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

VIDEO_LABEL = "<typ2pptx-video>"

# Allowed mismatch, in pt, between the metadata box and the poster picture.
MATCH_TOLERANCE_PT = 0.5

# mimetypes has no entry for some common video containers on every platform.
_VIDEO_MIME_TYPES = {
    ".mp4": "video/mp4",
    ".m4v": "video/mp4",
    ".mov": "video/quicktime",
    ".webm": "video/webm",
    ".wmv": "video/x-ms-wmv",
    ".avi": "video/x-msvideo",
}


@dataclass
class VideoPlacement:
    """One video on one physical page, in typst pt coordinates."""
    src: str
    page: int
    x: float
    y: float
    width: float
    height: float
    poster: Optional[str] = None

    @classmethod
    def from_metadata(cls, value: Dict[str, Any]) -> "VideoPlacement":
        return cls(
            src=str(value["src"]),
            poster=value.get("poster"),
            page=int(value["page"]),
            x=float(value["x"]),
            y=float(value["y"]),
            width=float(value["width"]),
            height=float(value["height"]),
        )


def query_videos(typ_path: str, root: Optional[str] = None) -> List[VideoPlacement]:
    """Query ``<typ2pptx-video>`` metadata from a .typ file with typst-py."""
    import typst as typst_py

    query_kwargs: Dict[str, Any] = {"field": "value"}
    if root is not None:
        query_kwargs["root"] = root
    result = typst_py.query(str(Path(typ_path).resolve()), VIDEO_LABEL, **query_kwargs)
    return [VideoPlacement.from_metadata(v) for v in json.loads(result)]


def _resolve(path: str, base_dir: Path) -> Path:
    # Typst paths are project-root relative; a leading "/" means the root too.
    return base_dir / path.lstrip("/")


def _video_mime_type(path: Path) -> str:
    return (_VIDEO_MIME_TYPES.get(path.suffix.lower())
            or mimetypes.guess_type(path.name)[0]
            or "video/unknown")


def _find_poster_picture(slide, box, tolerance_emu: int):
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        got = (shape.left, shape.top, shape.width, shape.height)
        if all(abs(a - b) <= tolerance_emu for a, b in zip(got, box)):
            return shape
    return None


def embed_videos(
    prs,
    videos: List[VideoPlacement],
    base_dir: str,
    emu_per_pt: float,
    verbose: bool = False,
) -> List[Dict[str, Any]]:
    """Replace poster pictures with embedded movies.

    Args:
        prs: Presentation produced by TypstSVGConverter (one slide per page)
        videos: Placements from ``query_videos``
        base_dir: Directory that ``src`` / ``poster`` paths are relative to
        emu_per_pt: Geometry scale used for the slides
            (``ConversionConfig.resolved_emu_per_px()``)
        verbose: Print each placement

    Returns:
        One dict per placement with ``status`` "embedded" or "skipped".
    """
    base = Path(base_dir)
    tolerance = round(MATCH_TOLERANCE_PT * emu_per_pt)
    report = []
    for v in videos:
        entry: Dict[str, Any] = {"src": v.src, "page": v.page}
        report.append(entry)
        if not 1 <= v.page <= len(prs.slides):
            entry["status"] = "skipped"
            entry["reason"] = f"page {v.page} outside 1..{len(prs.slides)}"
            print(f"Warning: video {v.src}: {entry['reason']}", file=sys.stderr)
            continue
        # Geometry is SVG pt * emu_per_pt from the top-left corner, exactly as
        # for every other shape. Deriving the scale from slide_width / page
        # width would be wrong: typst.ts rounds the SVG page size up to whole pt.
        box = tuple(round(c * emu_per_pt) for c in (v.x, v.y, v.width, v.height))
        entry["box_emu"] = box
        slide = prs.slides[v.page - 1]

        poster_pic = None
        if v.poster:
            poster_pic = _find_poster_picture(slide, box, tolerance)
            if poster_pic is None:
                entry["status"] = "skipped"
                entry["reason"] = "poster not rendered on this page (hidden subslide?)"
                if verbose:
                    print(f"  Video {v.src}: slide {v.page}: {entry['reason']}")
                continue

        src = _resolve(v.src, base)
        if not src.is_file():
            raise FileNotFoundError(f"Video file not found: {src}")
        poster = _resolve(v.poster, base) if v.poster else None
        if poster is not None and not poster.is_file():
            raise FileNotFoundError(f"Video poster not found: {poster}")

        mime = _video_mime_type(src)
        movie = slide.shapes.add_movie(
            str(src), *map(Emu, box),
            poster_frame_image=str(poster) if poster else None,
            mime_type=mime,
        )
        movie.name = f"Video {src.name}"
        if poster_pic is not None:
            # Take the poster's place in the z-order, then drop the static copy.
            poster_el = poster_pic._element
            poster_el.addprevious(movie._element)
            poster_el.getparent().remove(poster_el)

        entry.update(status="embedded", mime_type=mime)
        if verbose:
            print(f"  Video {v.src}: slide {v.page} at EMU {box} ({mime})")
    return report
