"""Tests for recovering font families from a PDF of the document."""
import pytest
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches

from typ2pptx.core.converter import ConversionConfig, TypstSVGConverter
from typ2pptx.core.fonts import PdfGlyph, apply_pdf_fonts, resolve_font
from typ2pptx.core.typst_svg_parser import (
    FontVariant, GlyphUse, PageData, TextSegment, TypstSVGData,
)

FAMILIES = ["Roboto", "Source Han Sans", "Source Han Sans JP", "Arial", "Times New Roman"]


class TestResolveFont:
    def test_family_without_spaces(self):
        font = resolve_font("SourceHanSansJP-Bold", FAMILIES)
        assert font.family == "Source Han Sans JP"
        assert font.style == "bold"

    def test_longest_family_prefix(self):
        assert resolve_font("ArialMT", FAMILIES).family == "Arial"
        assert resolve_font("TimesNewRomanPSMT", FAMILIES).family == "Times New Roman"

    @pytest.mark.parametrize("name, style", [
        ("Roboto-Regular", "regular"),
        ("Roboto-Bold", "bold"),
        ("Roboto-Italic", "italic"),
        ("Roboto-BoldItalic", "bolditalic"),
        ("Roboto-SemiBold", "bold"),
        ("Roboto-Medium", "regular"),
        ("SourceHanSansJP-DemiLight", "regular"),
    ])
    def test_style_from_name(self, name, style):
        assert resolve_font(name, FAMILIES).style == style

    def test_math_font(self):
        assert resolve_font("NewCMMath-Book", FAMILIES).style == "math"

    def test_mono_font(self):
        assert resolve_font("DejaVuSansMono-Bold", FAMILIES).style == "mono"

    def test_font_not_installed_has_no_family(self):
        font = resolve_font("LibertinusSerif-Regular", FAMILIES)
        assert font.family is None
        assert font.style == "regular"


def _segment(text, x, baseline, prefix, size=24.0):
    return TextSegment(
        text=text, x=x, y=baseline - size * 0.8, width=size * len(text), height=size,
        font_size=size, font_variant="mono", fill_color="#000000",
        glyph_uses=[GlyphUse(glyph_id=f"{prefix}{i:03}", x_offset=0, y_offset=0, path_data="")
                    for i in range(len(text))],
    )


def _svg_data(*segments):
    page = PageData(page_num=0, y_offset=0, width=960, height=540, text_segments=list(segments))
    prefixes = {s.glyph_uses[0].glyph_id[:5] for s in segments}
    variants = {p: FontVariant(prefix=p, glyph_count=1, uses_quadratic=True, style="mono")
                for p in prefixes}
    return TypstSVGData(viewbox_width=960, viewbox_height=540, pages=[page], font_variants=variants)


class TestApplyPdfFonts:
    def test_matches_segments_by_position(self):
        latin = _segment("Title", 44.8, 90.5, "gAAAA")
        cjk = _segment("研究", 120.0, 90.5, "gBBBB")
        data = _svg_data(latin, cjk)
        pdf = [[
            PdfGlyph("T", 44.84, 90.73, "Roboto-Regular"),
            PdfGlyph("研", 120.0, 90.73, "SourceHanSansJP-Bold"),
        ]]

        apply_pdf_fonts(data, pdf, FAMILIES)

        assert (latin.font_family, latin.font_variant) == ("Roboto", "regular")
        assert (cjk.font_family, cjk.font_variant) == ("Source Han Sans JP", "bold")
        assert data.font_variants["gBBBB"].style == "bold"

    def test_prefix_applies_to_unmatched_segments(self):
        first = _segment("Hello", 10, 100, "gAAAA")
        second = _segment("World", 10, 200, "gAAAA")
        apply_pdf_fonts(_svg_data(first, second), [[PdfGlyph("H", 10, 100, "Roboto-Bold")]], FAMILIES)
        assert second.font_family == "Roboto"
        assert second.font_variant == "bold"

    def test_leading_space_and_distant_glyphs(self):
        seg = _segment(" x", 10, 100, "gAAAA")
        far = PdfGlyph("x", 10, 150, "Roboto-Bold")  # another line
        near = PdfGlyph("x", 16, 100, "Roboto-Italic")  # after the space
        apply_pdf_fonts(_svg_data(seg), [[far, near]], FAMILIES)
        assert seg.font_variant == "italic"

    def test_no_match_keeps_heuristic(self):
        seg = _segment("Hello", 10, 100, "gAAAA")
        apply_pdf_fonts(_svg_data(seg), [[PdfGlyph("H", 300, 400, "Roboto-Bold")]], FAMILIES)
        assert seg.font_family is None
        assert seg.font_variant == "mono"


class TestRunFont:
    def _run(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        box = slide.shapes.add_textbox(0, 0, Inches(1), Inches(1))
        return box.text_frame.paragraphs[0].add_run()

    def test_family_sets_latin_and_east_asian(self):
        run = self._run()
        seg = _segment("研究", 0, 0, "gAAAA")
        seg.font_variant, seg.font_family = "bold", "Source Han Sans JP"
        TypstSVGConverter(ConversionConfig())._set_run_font(run, seg)

        rPr = run._r.rPr
        assert rPr.find(qn("a:latin")).get("typeface") == "Source Han Sans JP"
        assert rPr.find(qn("a:ea")).get("typeface") == "Source Han Sans JP"
        assert run.font.bold

    def test_without_family_falls_back(self):
        run = self._run()
        seg = _segment("x = 1", 0, 0, "gAAAA")
        TypstSVGConverter(ConversionConfig())._set_run_font(run, seg)
        assert run.font.name == "Consolas"
        assert run._r.rPr.find(qn("a:ea")) is None

    def test_math_keeps_cambria(self):
        run = self._run()
        seg = _segment("x", 0, 0, "gAAAA")
        seg.font_variant, seg.font_family = "math", "New Computer Modern Math"
        TypstSVGConverter(ConversionConfig())._set_run_font(run, seg)
        assert run.font.name == "Cambria Math"


@pytest.fixture(scope="module")
def runs(typ_sources_dir, typst_ts_cli, output_dir):
    """Text runs of chinese_test.typ converted from .typ input."""
    from typ2pptx.core.converter import convert_typst_to_pptx
    out = convert_typst_to_pptx(
        str(typ_sources_dir / "chinese_test.typ"),
        output_path=str(output_dir / "fonts_chinese_test.pptx"),
        typst_ts_cli=typst_ts_cli,
    )
    prs = Presentation(out)
    return [r for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
            for p in shape.text_frame.paragraphs for r in p.runs]


class TestConvertTypst:
    """End-to-end: .typ input recovers fonts from a PDF of the document."""

    def test_named_families_are_installed(self, runs):
        from typ2pptx.core.fonts import installed_families
        installed = set(installed_families())
        fallbacks = {"Arial", "Consolas", "Cambria Math"}
        for run in runs:
            assert run.font.name in installed | fallbacks, run.text

    def test_named_families_cover_east_asian_text(self, runs):
        for run in runs:
            ea = run._r.rPr.find(qn("a:ea"))
            if ea is not None:
                assert ea.get("typeface") == run.font.name
